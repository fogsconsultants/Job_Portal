from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = "JobPortal_Study_Guide.docx"
doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Inches(0.8)
section.left_margin = section.right_margin = Inches(0.85)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal'].font.size = Pt(10.5)
styles['Normal'].paragraph_format.space_after = Pt(6)
styles['Heading 1'].font.name = 'Aptos Display'
styles['Heading 1'].font.size = Pt(18)
styles['Heading 1'].font.color.rgb = RGBColor(24, 49, 46)
styles['Heading 1'].paragraph_format.space_before = Pt(16)
styles['Heading 1'].paragraph_format.space_after = Pt(7)
styles['Heading 2'].font.name = 'Aptos Display'
styles['Heading 2'].font.size = Pt(13)
styles['Heading 2'].font.color.rgb = RGBColor(181, 79, 59)
styles['Heading 2'].paragraph_format.space_before = Pt(10)
styles['Heading 2'].paragraph_format.space_after = Pt(5)

def title(text, size, color, after):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text); r.bold = True; r.font.name = 'Aptos Display'; r.font.size = Pt(size); r.font.color.rgb = color

def h(text, level=1): doc.add_heading(text, level=level)
def p(text, bold_prefix=None):
    para = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        run = para.add_run(bold_prefix); run.bold = True; para.add_run(text[len(bold_prefix):])
    else: para.add_run(text)
    return para
def bullets(items):
    for item in items: doc.add_paragraph(item, style='List Bullet')
def steps(items):
    for item in items: doc.add_paragraph(item, style='List Number')
def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)
def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.style='Table Grid'
    for i, text in enumerate(headers):
        cell=t.rows[0].cells[i]; cell.text=text; shade(cell,'18312E')
        for r in cell.paragraphs[0].runs: r.font.color.rgb=RGBColor(255,255,255); r.bold=True
    for row in rows:
        cells=t.add_row().cells
        for i, text in enumerate(row): cells[i].text=text
    doc.add_paragraph()

title('CareerCanvas', 28, RGBColor(24,49,46), 4)
title('Django Job Portal — Project Study Guide', 15, RGBColor(181,79,59), 16)
p('A practical guide to the completed JobPortal project: what it is, how its parts connect, and how to study or extend it.')

h('1. What this project belongs to')
p('CareerCanvas is a web application in the job-board / recruitment-platform category. It is built as a server-rendered Django application using the Model–View–Template (MVT) architectural pattern. It also demonstrates full-stack CRUD fundamentals: data is stored in SQLite, Django views process requests, and HTML templates show the result in the browser.')
table(['Layer', 'Used here', 'Responsibility'], [
['Model', 'jobs/models.py', 'Defines sectors, companies, jobs and applications.'],
['View', 'jobs/views.py', 'Fetches data, validates actions, redirects and renders pages.'],
['Template', 'templates/', 'Presents reusable HTML pages to the user.'],
['Static assets', 'static/css/site.css', 'Supplies the visual design and responsive layout.'],
['Routing', 'jobs/urls.py + config/urls.py', 'Maps browser URLs to view functions.'],
])

h('2. Project map')
table(['Path', 'Purpose'], [
['config/settings.py', 'Django configuration: installed apps, template directory, static files and database.'],
['config/urls.py', 'Top-level URL dispatcher; includes the jobs app routes.'],
['jobs/models.py', 'Database schema for Sectors, Company, Jobs and Applications.'],
['jobs/forms.py', 'JobForm controls the fields and widgets used when posting or editing a job.'],
['jobs/views.py', 'Function-based pages for browsing, authentication, job management and applications.'],
['jobs/urls.py', 'Named routes such as all_jobs, job_detail, add_job and apply_job.'],
['templates/base.html', 'Shared site shell: header, navigation, messages, content block and footer.'],
['templates/*.html', 'Individual pages that extend base.html.'],
['static/css/site.css', 'Single responsive design system used by every page.'],
['db.sqlite3', 'Local development database.'],
])

h('3. How a request works')
steps([
'A visitor opens a URL, for example /.',
'config/urls.py sends the request into jobs/urls.py.',
'jobs/urls.py chooses the allJobs view.',
'The view queries Jobs and Sectors through Django’s ORM.',
'The view sends the results to all-jobs.html.',
'all-jobs.html extends base.html, and site.css styles the completed page.',
])

h('4. The data model')
p('The model relationships are the heart of the platform: one sector can contain many jobs; one job can receive many applications; one user can make applications. A database constraint prevents the same user from applying to the same job twice.')
table(['Model', 'Important fields', 'Relationship'], [
['Sectors', 'name', 'One sector can have many jobs.'],
['Company', 'name, address, website', 'Available for employer information; ready to connect to jobs later.'],
['Jobs', 'title, sector, location, description, salary, post_date, is_active', 'Each job belongs to one sector.'],
['Applications', 'job, user, date', 'Connects an authenticated user to a job.'],
])

h('5. Shared template and styling')
p('base.html is the reusable layout. Every content page uses {% extends "base.html" %} and replaces only its content block. This avoids duplicated navigation, metadata, footer and message markup. The stylesheet uses CSS variables for the palette, responsive grids for job cards, and a mobile media query for small screens.')
bullets([
'The header adapts to authentication: signed-in users can post jobs and sign out; guests see sign-in and account-creation links.',
'Django messages appear in a single shared location after events such as registration or applying for a role.',
'Job cards, forms, authentication screens and the detail page share one visual language.',
'The page works on mobile because grid layouts collapse to one column below 700px.',
])

h('6. Main user flows')
table(['Flow', 'Route / view', 'What happens'], [
['Browse', '/ → allJobs', 'Shows active jobs and an optional sector filter.'],
['Inspect role', '/job-detail/<id>/ → jobDetail', 'Shows the description and application state.'],
['Register', '/register/ → signUp', 'Creates a Django user using their email as the username.'],
['Sign in', '/login/ → signIn', 'Authenticates the user and starts a session.'],
['Post job', '/add-job/ → addJob', 'Protected page that saves a validated JobForm.'],
['Apply', '/apply-job/<id>/ → applyJob', 'POST-only protected action that creates a unique application.'],
])

h('7. Run the project locally')
steps([
'Open a terminal in the JobPortal folder.',
'Activate a Python environment that has Django installed (or install the project requirements).',
'Run: python manage.py migrate',
'Run: python manage.py runserver',
'Open http://127.0.0.1:8000/ in a browser.',
'Use /admin/ after creating an admin user with python manage.py createsuperuser to add sectors and sample jobs.',
])

h('8. Study order')
steps([
'Start with jobs/models.py to understand what data exists.',
'Read jobs/urls.py and then jobs/views.py to follow each browser request.',
'Compare a view’s context dictionary with the matching template variables.',
'Open base.html and then all-jobs.html to see template inheritance in practice.',
'Read static/css/site.css section by section and use browser developer tools to test a small style change.',
'Create a sector and job in the admin, then trace it from the database to the job card.',
'Register a test user and submit an application to observe sessions, POST requests and messages.',
])

h('9. Good next improvements')
bullets([
'Add a Company foreign key to Jobs and show the company on cards and detail pages.',
'Let employers see and manage applications for their own job posts.',
'Add search, pagination, salary range and location filtering.',
'Upload resumes safely with file validation and media-file configuration.',
'Use Django permissions so only a job’s owner can edit it.',
'Move secrets and DEBUG settings to environment variables before deployment.',
'Add tests for authentication, job visibility and duplicate-application behaviour.',
])

h('10. Key terms to remember')
table(['Term', 'Meaning in this project'], [
['Django ORM', 'Python interface used to query and save database records without writing SQL for everyday actions.'],
['Migration', 'Versioned instruction that changes the database schema, such as the unique application constraint.'],
['CSRF token', 'Security token included in each form that protects POST requests.'],
['Session', 'Server-side mechanism Django uses to remember a signed-in user.'],
['Template inheritance', 'A child page extends base.html rather than repeating the entire HTML document.'],
['Static files', 'CSS, JavaScript and images served separately from template HTML.'],
])

p('Study tip: make one small change at a time, refresh the browser, then identify which model, view, template or stylesheet was responsible. That feedback loop is the fastest way to become comfortable with Django.', bold_prefix='Study tip: ')
doc.save(OUT)
